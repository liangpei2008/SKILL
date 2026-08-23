from __future__ import annotations

import argparse
import copy
import hashlib
import json
import re
import shutil
from io import BytesIO
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
CN_NUM = "一二三四五六七八九十"


HEADING_RULES = {
    "Heading 1": ("黑体", 18.0, True, WD_ALIGN_PARAGRAPH.JUSTIFY, 2.0, 6.0, 6.0, 0.0),
    "Heading 2": ("楷体", 18.0, True, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.5, None, None, 0.0),
    "Heading 3": ("仿宋", 16.0, True, WD_ALIGN_PARAGRAPH.LEFT, 1.5, None, None, 0.0),
    "Heading 4": ("仿宋", 15.0, True, WD_ALIGN_PARAGRAPH.LEFT, 1.5, None, None, 0.0),
}
NORMAL_RULE = ("仿宋", 14.0, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.5, None, None, 0.97)
TABLE_RULE = ("仿宋", 12.0, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.0, None, None, 0.0)
CAPTION_RULE = ("仿宋", 14.0, True, WD_ALIGN_PARAGRAPH.CENTER, 1.5, None, None, 0.0)
CAPTION_PATTERN = re.compile(r"^[图表]\s*[0-9一二三四五六七八九十]+(?:[.．、：:\-－])?\s+.+")
CONTENT_WIDTH_CM = 14.66
TABLE_PROFILES = [
    {
        "contains": ["序号", "考核指标", "完成情况", "是否已完成"],
        "widths_cm": [0.94, 2.89, 4.17, 1.74, 2.55, 2.73],
        "border": {"style": "none", "size": 0, "color": "auto"},
    },
    {
        "contains": ["指标", "市财政资金", "自筹资金", "合计"],
        "widths_cm": [4.82, 4.06, 3.08, 3.08],
        "border": {"style": "single", "size": 4, "color": "auto"},
    },
    {
        "contains": ["序号", "财政拨款资金支出科目", "计划支出金额", "已支付资金"],
        "widths_cm": [2.35, 5.93, 3.68, 3.06],
        "border": {"style": "single", "size": 6, "color": "auto"},
    },
]
CONTENT_SUFFIXES = [".docx", ".pdf", ".png", ".jpg", ".jpeg"]


def normalize_text(text: str) -> str:
    value = "".join((text or "").split())
    value = re.sub(r"^[一二三四五六七八九十]+、", "", value)
    value = re.sub(r"^（[一二三四五六七八九十]+）", "", value)
    value = re.sub(r"^\d+(?:\.\d+)*[\.．、]?", "", value)
    return value


def clean_heading_name(text: str) -> str:
    value = (text or "").strip()
    patterns = [
        r"^[一二三四五六七八九十]+、\s*",
        r"^（[一二三四五六七八九十]+）\s*",
        r"^\d+(?:\.\d+)*[\.．、]?\s*",
    ]
    changed = True
    while changed:
        changed = False
        for pattern in patterns:
            new_value = re.sub(pattern, "", value)
            if new_value != value:
                value = new_value
                changed = True
    return value.strip()


def clear_body(doc: Document) -> None:
    body = doc._body._element
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def insert_before_sectpr(doc: Document, element) -> None:
    body = doc._body._element
    sect_pr = body.sectPr
    if sect_pr is None:
        body.append(element)
    else:
        body.insert(body.index(sect_pr), element)


def iter_body_elements(doc: Document) -> Iterable:
    for element in doc._body._element:
        if element.tag != qn("w:sectPr"):
            yield element


def element_text(element) -> str:
    return "".join(t.text or "" for t in element.iter(qn("w:t"))).strip()


def element_style_id(element) -> str:
    ppr = element.find(qn("w:pPr"))
    if ppr is None:
        return ""
    pstyle = ppr.find(qn("w:pStyle"))
    if pstyle is None:
        return ""
    return pstyle.get(qn("w:val")) or ""


def heading_level_from_style_id(style_id: str) -> int | None:
    mapping = {
        "Heading1": 1,
        "Heading2": 2,
        "Heading3": 3,
        "Heading4": 4,
        # Common localized built-in ids observed in Chinese Word files.
        "2": 1,
        "3": 2,
        "4": 3,
        "5": 4,
    }
    return mapping.get(style_id)


def set_style_id(element, style_id: str) -> None:
    ppr = element.find(qn("w:pPr"))
    if ppr is None:
        ppr = OxmlElement("w:pPr")
        element.insert(0, ppr)
    pstyle = ppr.find(qn("w:pStyle"))
    if pstyle is None:
        pstyle = OxmlElement("w:pStyle")
        ppr.insert(0, pstyle)
    pstyle.set(qn("w:val"), style_id)


def heading_style_id(doc: Document, level: int) -> str:
    """Return the actual style id used by the destination document.

    Chinese Word files often localize built-in Heading styles to numeric ids
    such as 2/3/4/5. Writing literal Heading2 into pStyle can make Word fall
    back to Normal if that style id is not defined in the destination package.
    """
    level = max(1, min(level, 4))
    return doc.styles[f"Heading {level}"].style_id


def copy_relationships(src_doc: Document, dst_doc: Document, element) -> None:
    for attr in ("embed", "link", "id"):
        qattr = f"{{{REL_NS}}}{attr}"
        for node in element.xpath(f".//*[@r:{attr}]"):
            old_rid = node.get(qattr)
            if not old_rid or old_rid not in src_doc.part.rels:
                continue
            rel = src_doc.part.rels[old_rid]
            if rel.is_external:
                new_rid = dst_doc.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
            elif rel.reltype.endswith("/image"):
                new_rid, _image = dst_doc.part.get_or_add_image(BytesIO(rel.target_part.blob))
            else:
                continue
            node.set(qattr, new_rid)


def first_nonempty_heading(elements: list) -> tuple[int, int | None, str] | None:
    for idx, element in enumerate(elements):
        text = element_text(element)
        if not text:
            continue
        if element.tag != qn("w:p"):
            return idx, None, text
        return idx, heading_level_from_style_id(element_style_id(element)), text
    return None


def add_heading(doc: Document, text: str, level: int) -> None:
    para = doc.add_paragraph(style=f"Heading {max(1, min(level, 4))}")
    para.add_run(text)


def add_h1(doc: Document, text: str) -> None:
    add_heading(doc, text, 1)


def append_chapter(doc: Document, title: str, source: Path) -> dict:
    add_h1(doc, title)
    src = Document(str(source))
    elements = list(iter_body_elements(src))
    first = first_nonempty_heading(elements)
    skip_until = 0
    demote = 0
    if first:
        idx, level, text = first
        if level == 1 and normalize_text(text) == normalize_text(title):
            skip_until = idx + 1
        elif level == 1:
            demote = 1

    copied = 0
    headings = []
    for element in elements[skip_until:]:
        cloned = copy.deepcopy(element)
        if cloned.tag == qn("w:p"):
            level = heading_level_from_style_id(element_style_id(cloned))
            if level:
                new_level = min(level + demote, 4)
                set_style_id(cloned, heading_style_id(doc, new_level))
                text = element_text(cloned)
                headings.append({"level": new_level, "text": text})
        copy_relationships(src, doc, cloned)
        insert_before_sectpr(doc, cloned)
        copied += 1
    return {
        "title": title,
        "source": str(source),
        "skipped_leading_heading": bool(skip_until),
        "demoted_levels": demote,
        "copied_elements": copied,
        "headings": headings,
    }


def append_docx_under_heading(doc: Document, title: str, level: int, source: Path) -> dict:
    add_heading(doc, title, level)
    src = Document(str(source))
    elements = list(iter_body_elements(src))
    first = first_nonempty_heading(elements)
    skip_until = 0
    base_level = 0
    if first:
        idx, first_level, text = first
        if first_level and normalize_text(text) == normalize_text(title):
            skip_until = idx + 1
            base_level = first_level

    copied = 0
    headings = []
    for element in elements[skip_until:]:
        cloned = copy.deepcopy(element)
        if cloned.tag == qn("w:p"):
            source_level = heading_level_from_style_id(element_style_id(cloned))
            if source_level:
                if base_level:
                    new_level = level + max(1, source_level - base_level)
                else:
                    new_level = level + source_level
                new_level = min(new_level, 4)
                set_style_id(cloned, heading_style_id(doc, new_level))
                headings.append({"level": new_level, "text": element_text(cloned)})
        copy_relationships(src, doc, cloned)
        insert_before_sectpr(doc, cloned)
        copied += 1
    return {
        "title": title,
        "level": level,
        "source": str(source),
        "type": "docx",
        "skipped_leading_heading": bool(skip_until),
        "copied_elements": copied,
        "headings": headings,
    }


def set_run_font(run, font_name: str, size_pt: float, bold: bool | None) -> None:
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("eastAsia", "ascii", "hAnsi", "cs"):
        rfonts.set(qn(f"w:{attr}"), font_name)


def apply_style_rule(style, rule: tuple) -> None:
    font_name, size, bold, align, line_spacing, before, after, first_indent = rule
    style.font.name = font_name
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("eastAsia", "ascii", "hAnsi", "cs"):
        rfonts.set(qn(f"w:{attr}"), font_name)
    fmt = style.paragraph_format
    fmt.alignment = align
    fmt.line_spacing = line_spacing
    fmt.first_line_indent = Cm(first_indent)
    if before is not None:
        fmt.space_before = Pt(before)
    if after is not None:
        fmt.space_after = Pt(after)


def apply_paragraph_rule(paragraph, rule: tuple) -> None:
    font_name, size, bold, align, line_spacing, before, after, first_indent = rule
    fmt = paragraph.paragraph_format
    fmt.alignment = align
    fmt.line_spacing = line_spacing
    fmt.first_line_indent = Cm(first_indent)
    if before is not None:
        fmt.space_before = Pt(before)
    if after is not None:
        fmt.space_after = Pt(after)
    for run in paragraph.runs:
        if run._element.xpath(".//w:drawing"):
            continue
        set_run_font(run, font_name, size, bold)


def cm_to_twips(value: float) -> int:
    return int(round(value * 567))


def set_or_append(parent, tag: str):
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    return child


def set_table_width(table, width_cm: float) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = set_or_append(tbl_pr, "w:tblW")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(cm_to_twips(width_cm)))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is not None:
        tbl_pr.remove(tbl_ind)


def set_table_grid(table, widths_cm: list[float]) -> None:
    tbl_grid = table._tbl.tblGrid
    if tbl_grid is None:
        tbl_grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, tbl_grid)
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths_cm:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(cm_to_twips(float(width))))
        tbl_grid.append(col)


def set_cell_width(cell, width_cm: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = set_or_append(tc_pr, "w:tcW")
    tc_w.set(qn("w:type"), "dxa")
    tc_w.set(qn("w:w"), str(cm_to_twips(width_cm)))
    cell.width = Cm(width_cm)


def set_cell_margins(cell, top: int = 80, bottom: int = 80, left: int = 120, right: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = set_or_append(tc_pr, "w:tcMar")
    for side, value in {"top": top, "bottom": bottom, "left": left, "right": right}.items():
        node = set_or_append(tc_mar, f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, style: str = "single", size: int = 4, color: str = "auto") -> None:
    tbl_pr = table._tbl.tblPr
    borders = set_or_append(tbl_pr, "w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = set_or_append(borders, f"w:{edge}")
        node.set(qn("w:val"), style)
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def clear_table_shading(table) -> None:
    for row in table.rows:
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            shd = tc_pr.find(qn("w:shd"))
            if shd is not None:
                tc_pr.remove(shd)


def table_text(table) -> str:
    return " / ".join(cell.text.strip() for cell in table.rows[0].cells) if table.rows else ""


def table_profile(table, manifest_profiles: list[dict] | None = None) -> dict:
    header = table_text(table)
    profiles = manifest_profiles or []
    for profile in profiles + TABLE_PROFILES:
        if all(token in header for token in profile.get("contains", [])):
            return profile
    cols = len(table.columns)
    width = CONTENT_WIDTH_CM / cols if cols else CONTENT_WIDTH_CM
    return {"widths_cm": [width] * cols, "border": {"style": "single", "size": 4, "color": "auto"}}


def apply_table_geometry(table, manifest_profiles: list[dict] | None = None) -> None:
    profile = table_profile(table, manifest_profiles)
    widths = profile.get("widths_cm") or []
    border = profile.get("border", {"style": "single", "size": 4, "color": "auto"})
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if widths:
        widths = [float(w) for w in widths]
        table_width = sum(float(w) for w in widths)
        set_table_width(table, table_width)
        set_table_grid(table, widths)
    set_table_borders(
        table,
        style=border.get("style", "single"),
        size=int(border.get("size", 4)),
        color=border.get("color", "auto"),
    )
    clear_table_shading(table)
    for row_idx, row in enumerate(table.rows):
        if row_idx == 0:
            row._tr.get_or_add_trPr()
            tbl_header = row._tr.trPr.find(qn("w:tblHeader"))
            if tbl_header is None:
                row._tr.trPr.append(OxmlElement("w:tblHeader"))
        for idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if idx < len(widths):
                set_cell_width(cell, float(widths[idx]))
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                apply_paragraph_rule(paragraph, TABLE_RULE)


def is_caption(text: str) -> bool:
    return bool(CAPTION_PATTERN.match((text or "").strip()))


def normalize_caption_prefix(paragraph) -> None:
    """Remove leading spaces/tabs before figure/table captions.

    A caption such as " 图1 ..." visually shifts the caption even when the
    paragraph is centered in Word. Rewriting the caption text is acceptable
    here because formatting is applied immediately afterward.
    """
    text = paragraph.text
    stripped = text.lstrip(" \t\u3000")
    if text == stripped or not is_caption(stripped):
        return
    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = stripped
    else:
        paragraph.add_run(stripped)


def rule_for_paragraph(paragraph) -> tuple:
    name = paragraph.style.name if paragraph.style is not None else ""
    if name in HEADING_RULES:
        return HEADING_RULES[name]
    if is_caption(paragraph.text):
        return CAPTION_RULE
    return NORMAL_RULE


def apply_format(doc: Document, image_width_cm: float = 16.13, table_profiles: list[dict] | None = None) -> None:
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.17)
        section.right_margin = Cm(3.17)
        section.header_distance = Cm(1.5)
        section.footer_distance = Cm(1.75)

    for style_name, rule in HEADING_RULES.items():
        apply_style_rule(doc.styles[style_name], rule)
    apply_style_rule(doc.styles["Normal"], NORMAL_RULE)

    for paragraph in doc.paragraphs:
        if is_caption(paragraph.text):
            normalize_caption_prefix(paragraph)
        apply_paragraph_rule(paragraph, rule_for_paragraph(paragraph))
    for table in doc.tables:
        apply_table_geometry(table, table_profiles)
    for paragraph in doc.paragraphs:
        if paragraph._p.xpath(".//w:drawing"):
            paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.line_spacing = 1.5
    for shape in doc.inline_shapes:
        if shape.width:
            ratio = Cm(image_width_cm) / shape.width
            shape.width = Cm(image_width_cm)
            shape.height = int(shape.height * ratio)


def strip_heading_prefix(text: str) -> str:
    value = text.strip()
    patterns = [
        r"^[一二三四五六七八九十]+、\s*",
        r"^（[一二三四五六七八九十]+）\s*",
        r"^[（(]\d+[）)]\s*",
        r"^\d+(?:\.\d+)*[\.．、]?\s*",
        r"^任务[一二三四五六七八九十]+[：:、]\s*",
    ]
    changed = True
    while changed:
        changed = False
        for pattern in patterns:
            new_value = re.sub(pattern, "", value)
            if new_value != value:
                value = new_value
                changed = True
    return value


def cn_index(idx: int) -> str:
    if idx <= 10:
        return CN_NUM[idx - 1]
    if idx < 20:
        return "十" + (CN_NUM[idx - 11] if idx > 10 else "")
    tens, ones = divmod(idx, 10)
    return CN_NUM[tens - 1] + "十" + (CN_NUM[ones - 1] if ones else "")


def heading_level_from_text(text: str) -> int | None:
    value = (text or "").strip()
    if re.match(r"^[一二三四五六七八九十]+、\S+", value):
        return 1
    if re.match(r"^（[一二三四五六七八九十]+）\S+", value):
        return 2
    if re.match(r"^\d+\.\d+\.\S+", value):
        return 4
    if re.match(r"^\d+[\.．、]\S+", value):
        return 3
    return None


def paragraph_heading_level(paragraph) -> int | None:
    if paragraph.style is not None:
        level = heading_level_from_style_id(paragraph.style.style_id)
        if level:
            return level
        if paragraph.style.name.startswith("Heading "):
            try:
                return int(paragraph.style.name.split()[-1])
            except ValueError:
                pass
    return heading_level_from_text(paragraph.text)


def infer_outline_from_master(master: Path) -> list[dict]:
    doc = Document(str(master))
    nodes: list[dict] = []
    seen_paths = set()
    stack: list[dict] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        level = paragraph_heading_level(paragraph)
        if not level or level > 4:
            continue
        name = clean_heading_name(text)
        if not name or normalize_text(name) in {"目录", "正文目录"}:
            continue
        while stack and stack[-1]["level"] >= level:
            stack.pop()
        parent_path = tuple(node["name"] for node in stack)
        key = parent_path + (normalize_text(name),)
        if key in seen_paths:
            continue
        seen_paths.add(key)
        node = {"title": text, "name": name, "level": level, "children": []}
        if stack:
            stack[-1]["children"].append(node)
        else:
            nodes.append(node)
        stack.append(node)
    return nodes


def renumber_outline(nodes: list[dict]) -> list[dict]:
    def walk(items: list[dict], counters: list[int]) -> None:
        for idx, node in enumerate(items, 1):
            level = node["level"]
            if level == 1:
                node["title"] = f"{cn_index(idx)}、{node['name']}"
                walk(node["children"], [idx, 0, 0, 0])
            elif level == 2:
                node["title"] = f"（{cn_index(idx)}）{node['name']}"
                next_counters = counters[:]
                next_counters[1] = idx
                walk(node["children"], next_counters)
            elif level == 3:
                node["title"] = f"{idx}.{node['name']}"
                next_counters = counters[:]
                next_counters[2] = idx
                walk(node["children"], next_counters)
            elif level == 4:
                prefix = f"{counters[2]}.{idx}" if len(counters) > 2 and counters[2] else str(idx)
                node["title"] = f"{prefix}.{node['name']}"
                walk(node["children"], counters)

    walk(nodes, [0, 0, 0, 0])
    return nodes


def outline_top_headings(nodes: list[dict]) -> list[str]:
    return [node["title"] for node in nodes if node["level"] == 1]


def infer_top_headings_from_master(master: Path) -> list[str]:
    return outline_top_headings(renumber_outline(infer_outline_from_master(master)))


def candidate_files(chapter_dir: Path, excluded: set[Path], recursive: bool = True) -> list[Path]:
    files = []
    iterator = chapter_dir.rglob("*") if recursive else chapter_dir.glob("*")
    for path in iterator:
        if not path.is_file() or path.suffix.lower() not in CONTENT_SUFFIXES:
            continue
        if path.name.startswith("~$"):
            continue
        if path.resolve() in excluded:
            continue
        if any(token in path.stem for token in ("备份", "backup", "格式优化")):
            continue
        files.append(path)
    return files


def candidate_docx_files(chapter_dir: Path, excluded: set[Path]) -> list[Path]:
    return [path for path in candidate_files(chapter_dir, excluded) if path.suffix.lower() == ".docx"]


def match_content_file(title: str, candidates: list[Path], suffixes: list[str] | None = None) -> Path | None:
    target = normalize_text(clean_heading_name(title))
    allowed = {suffix.lower() for suffix in suffixes} if suffixes else None
    scoped = [path for path in candidates if allowed is None or path.suffix.lower() in allowed]
    exact = [path for path in scoped if clean_heading_name(path.stem) == clean_heading_name(title)]
    if len(exact) == 1:
        return exact[0]
    normalized = [path for path in scoped if normalize_text(clean_heading_name(path.stem)) == target]
    if len(normalized) == 1:
        return normalized[0]
    return None


def match_chapter_file(title: str, candidates: list[Path]) -> Path | None:
    return match_content_file(title, candidates, suffixes=[".docx"])


def infer_chapters_from_titles(top_headings: list[str], chapter_dir: Path, master: Path, output: Path) -> list[dict]:
    excluded = {master.resolve(), output.resolve()}
    candidates = candidate_docx_files(chapter_dir, excluded)
    chapters = []
    missing = []
    for title in top_headings:
        source = match_chapter_file(title, candidates)
        if source is None:
            missing.append(clean_heading_name(title))
            continue
        chapters.append({"title": title, "source": str(source), "auto_matched": True})
    if missing:
        available = [path.name for path in candidates]
        raise FileNotFoundError(f"Cannot auto-match chapter files for headings: {missing}. Available docx files: {available}")
    return chapters


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def render_pdf_pages(pdf: Path, cache_root: Path, dpi: int = 160) -> list[Path]:
    pdf_hash = file_sha256(pdf)
    cache_dir = cache_root / "pdf-pages" / f"{pdf.stem}-{pdf_hash[:12]}"
    manifest = cache_dir / "manifest.json"
    if manifest.exists():
        data = json.loads(manifest.read_text(encoding="utf-8"))
        pages = [cache_dir / item for item in data.get("pages", [])]
        if pages and all(page.exists() for page in pages):
            return pages
    cache_dir.mkdir(parents=True, exist_ok=True)
    pages = []
    try:
        import fitz  # PyMuPDF

        doc = fitz.open(str(pdf))
        zoom = dpi / 72
        matrix = fitz.Matrix(zoom, zoom)
        for index, page in enumerate(doc, 1):
            image_path = cache_dir / f"page-{index:04d}.png"
            pix = page.get_pixmap(matrix=matrix, alpha=False)
            pix.save(str(image_path))
            pages.append(image_path)
    except ModuleNotFoundError:
        try:
            import pypdfium2 as pdfium
        except ModuleNotFoundError as exc:
            raise RuntimeError("PDF rendering requires PyMuPDF (fitz) or pypdfium2. Use the bundled Codex Python runtime or install one of these packages before binding PDF attachments.") from exc

        doc = pdfium.PdfDocument(str(pdf))
        scale = dpi / 72
        for index in range(len(doc)):
            image_path = cache_dir / f"page-{index + 1:04d}.png"
            page = doc[index]
            bitmap = page.render(scale=scale)
            pil_image = bitmap.to_pil()
            pil_image.save(image_path)
            pages.append(image_path)
    manifest.write_text(json.dumps({"source": str(pdf), "sha256": pdf_hash, "pages": [p.name for p in pages]}, ensure_ascii=False, indent=2), encoding="utf-8")
    return pages


def append_image(doc: Document, image_path: Path, width_cm: float = 16.13) -> None:
    para = doc.add_paragraph()
    para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.first_line_indent = Cm(0)
    run = para.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))


def append_pdf_under_heading(doc: Document, title: str, level: int, source: Path, cache_root: Path, image_width_cm: float) -> dict:
    add_heading(doc, title, level)
    pages = render_pdf_pages(source, cache_root)
    for image_path in pages:
        append_image(doc, image_path, image_width_cm)
    return {"title": title, "level": level, "source": str(source), "type": "pdf", "pages": len(pages)}


def append_image_under_heading(doc: Document, title: str, level: int, source: Path, image_width_cm: float) -> dict:
    add_heading(doc, title, level)
    append_image(doc, source, image_width_cm)
    return {"title": title, "level": level, "source": str(source), "type": "image", "pages": 1}


def append_content_node(doc: Document, node: dict, source: Path, cache_root: Path, image_width_cm: float) -> dict:
    suffix = source.suffix.lower()
    if suffix == ".docx":
        return append_docx_under_heading(doc, node["title"], node["level"], source)
    if suffix == ".pdf":
        return append_pdf_under_heading(doc, node["title"], node["level"], source, cache_root, image_width_cm)
    if suffix in {".png", ".jpg", ".jpeg"}:
        return append_image_under_heading(doc, node["title"], node["level"], source, image_width_cm)
    raise ValueError(f"Unsupported content file type: {source}")


def append_outline_nodes(doc: Document, nodes: list[dict], candidates: list[Path], cache_root: Path, image_width_cm: float, parent_has_source: bool = False) -> list[dict]:
    reports = []
    for node in nodes:
        suffixes = [".docx"] if node["level"] == 1 else CONTENT_SUFFIXES
        source = match_content_file(node["title"], candidates, suffixes=suffixes)
        if source:
            if node["level"] == 1 and source.suffix.lower() == ".docx":
                reports.append(append_chapter(doc, node["title"], source))
            else:
                reports.append(append_content_node(doc, node, source, cache_root, image_width_cm))
        elif node["children"]:
            add_heading(doc, node["title"], node["level"])
            reports.append({"title": node["title"], "level": node["level"], "type": "directory"})
            reports.extend(append_outline_nodes(doc, node["children"], candidates, cache_root, image_width_cm, parent_has_source=False))
        else:
            available = [path.name for path in candidates]
            raise FileNotFoundError(f"Cannot auto-match content file for catalogue node: {node['title']}. Available files: {available}")
    return reports


def renumber_headings(doc: Document, top_headings: list[str]) -> None:
    h1_idx = -1
    h2_idx = 0
    h3_idx = 0
    h4_idx = 0
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        name = paragraph.style.name if paragraph.style is not None else ""
        if name == "Heading 1":
            h1_idx += 1
            h2_idx = h3_idx = h4_idx = 0
            new_text = top_headings[h1_idx] if h1_idx < len(top_headings) else text
        elif name == "Heading 2":
            h2_idx += 1
            h3_idx = h4_idx = 0
            new_text = f"（{cn_index(h2_idx)}）{strip_heading_prefix(text)}"
        elif name == "Heading 3":
            h3_idx += 1
            h4_idx = 0
            new_text = f"{h3_idx}.{strip_heading_prefix(text)}"
        elif name == "Heading 4":
            h4_idx += 1
            new_text = f"{h3_idx}.{h4_idx}.{strip_heading_prefix(text)}" if h3_idx else f"{h4_idx}.{strip_heading_prefix(text)}"
        else:
            continue
        for run in paragraph.runs:
            run.text = ""
        if paragraph.runs:
            paragraph.runs[0].text = new_text
        else:
            paragraph.add_run(new_text)


def is_empty_body_paragraph(element) -> bool:
    return element.tag == qn("w:p") and not element_text(element) and not element.xpath(".//w:drawing")


def clean_empty_body_paragraphs(doc: Document) -> int:
    removed = 0
    body = doc._body._element
    for element in list(body):
        if element.tag == qn("w:sectPr"):
            continue
        if is_empty_body_paragraph(element):
            body.remove(element)
            removed += 1
    return removed


def font_name_from_run(run) -> str | None:
    rpr = run._element.rPr
    if rpr is None or rpr.rFonts is None:
        return None
    return rpr.rFonts.get(qn("w:eastAsia")) or rpr.rFonts.get(qn("w:ascii")) or rpr.rFonts.get(qn("w:hAnsi"))


def pt_value(length) -> float | None:
    return round(length.pt, 2) if length is not None else None


def twips_to_cm(value: str | None) -> float | None:
    if not value:
        return None
    try:
        return round(int(value) / 567, 2)
    except ValueError:
        return None


def table_width_cm(table) -> float | None:
    tbl_w = table._tbl.tblPr.find(qn("w:tblW")) if table._tbl.tblPr is not None else None
    return twips_to_cm(tbl_w.get(qn("w:w")) if tbl_w is not None else None)


def first_row_widths_cm(table) -> list[float | None]:
    if not table.rows:
        return []
    widths = []
    for cell in table.rows[0].cells:
        tc_w = cell._tc.tcPr.find(qn("w:tcW")) if cell._tc.tcPr is not None else None
        widths.append(twips_to_cm(tc_w.get(qn("w:w")) if tc_w is not None else None))
    return widths


def table_border_values(table) -> dict:
    borders = table._tbl.tblPr.find(qn("w:tblBorders")) if table._tbl.tblPr is not None else None
    result = {}
    if borders is None:
        return result
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is not None:
            result[edge] = {
                "style": node.get(qn("w:val")),
                "size": node.get(qn("w:sz")),
                "color": node.get(qn("w:color")),
            }
    return result


def compare_widths(actual: list[float | None], expected: list[float], tolerance: float = 0.08) -> bool:
    if len(actual) != len(expected):
        return False
    for actual_value, expected_value in zip(actual, expected):
        if actual_value is None or abs(actual_value - float(expected_value)) > tolerance:
            return False
    return True


def audit_table_format(table, table_idx: int, violations: list, manifest_profiles: list[dict] | None = None) -> dict:
    profile = table_profile(table, manifest_profiles)
    expected_widths = [float(w) for w in profile.get("widths_cm", [])]
    expected_border = profile.get("border", {"style": "single", "size": 4, "color": "auto"})
    actual_widths = first_row_widths_cm(table)
    actual_border = table_border_values(table)
    report = {
        "table": table_idx,
        "header": table_text(table),
        "alignment": str(table.alignment),
        "width_cm": table_width_cm(table),
        "first_row_widths_cm": actual_widths,
        "expected_widths_cm": expected_widths,
        "borders": actual_border,
    }
    if table.alignment != WD_TABLE_ALIGNMENT.CENTER:
        violations.append({"issue": "table_alignment", "table": table_idx, "actual": str(table.alignment), "expected": "CENTER"})
    if expected_widths and not compare_widths(actual_widths, expected_widths):
        violations.append({"issue": "table_column_width", "table": table_idx, "actual": actual_widths, "expected": expected_widths})
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        actual = actual_border.get(edge, {})
        if actual.get("style") != expected_border.get("style", "single") or str(actual.get("size")) != str(expected_border.get("size", 4)):
            violations.append({
                "issue": "table_border",
                "table": table_idx,
                "edge": edge,
                "actual": actual,
                "expected": expected_border,
            })
            break
    for row_idx, row in enumerate(table.rows, 1):
        for col_idx, cell in enumerate(row.cells, 1):
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    if not run.text.strip():
                        continue
                    if (
                        font_name_from_run(run) != TABLE_RULE[0]
                        or pt_value(run.font.size) != TABLE_RULE[1]
                        or bool(run.bold) != TABLE_RULE[2]
                    ):
                        violations.append({
                            "issue": "table_run_format",
                            "table": table_idx,
                            "row": row_idx,
                            "col": col_idx,
                            "font": font_name_from_run(run),
                            "size": pt_value(run.font.size),
                            "bold": run.bold,
                            "expected_font": TABLE_RULE[0],
                            "expected_size": TABLE_RULE[1],
                            "expected_bold": TABLE_RULE[2],
                            "text": run.text[:80],
                        })
                        return report
    return report


def audit(doc: Document, top_headings: list[str], table_profiles: list[dict] | None = None) -> dict:
    headings = []
    violations = []
    image_paragraphs = []
    caption_paragraphs = []
    for idx, paragraph in enumerate(doc.paragraphs, 1):
        raw_text = paragraph.text
        text = raw_text.strip()
        name = paragraph.style.name if paragraph.style is not None else ""
        if name.startswith("Heading"):
            headings.append({"paragraph": idx, "style": name, "text": text})
        if is_caption(raw_text):
            has_leading_space = raw_text != raw_text.lstrip(" \t\u3000")
            caption_paragraphs.append({
                "paragraph": idx,
                "text": text,
                "alignment": str(paragraph.paragraph_format.alignment),
                "has_leading_space": has_leading_space,
            })
            if has_leading_space:
                violations.append({"paragraph": idx, "issue": "caption_leading_space", "text": raw_text[:80]})
            if paragraph.paragraph_format.alignment != WD_ALIGN_PARAGRAPH.CENTER:
                violations.append({"paragraph": idx, "issue": "caption_alignment", "actual": str(paragraph.paragraph_format.alignment), "expected": "CENTER"})
        if paragraph._p.xpath(".//w:drawing"):
            image_paragraphs.append({"paragraph": idx, "alignment": str(paragraph.paragraph_format.alignment)})
            if paragraph.paragraph_format.alignment != WD_ALIGN_PARAGRAPH.CENTER:
                violations.append({"paragraph": idx, "issue": "image_paragraph_alignment", "actual": str(paragraph.paragraph_format.alignment), "expected": "CENTER"})
        if not text:
            continue
        rule = rule_for_paragraph(paragraph)
        expected_font, expected_size, expected_bold = rule[0], rule[1], rule[2]
        for run in paragraph.runs:
            if not run.text.strip() or run._element.xpath(".//w:drawing"):
                continue
            if (
                font_name_from_run(run) != expected_font
                or pt_value(run.font.size) != expected_size
                or (expected_bold is not None and bool(run.bold) != expected_bold)
            ):
                violations.append({
                    "paragraph": idx,
                    "issue": "run_format",
                    "font": font_name_from_run(run),
                    "size": pt_value(run.font.size),
                    "bold": run.bold,
                    "expected_font": expected_font,
                    "expected_size": expected_size,
                    "expected_bold": expected_bold,
                    "text": text[:80],
                })
                break
    h1 = [h["text"] for h in headings if h["style"] == "Heading 1"]
    if h1 != top_headings:
        violations.append({"issue": "heading1_mismatch", "actual": h1, "expected": top_headings})
    table_reports = [audit_table_format(table, idx, violations, table_profiles) for idx, table in enumerate(doc.tables, 1)]
    return {
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "inline_shapes": len(doc.inline_shapes),
        "headings": headings,
        "heading1": h1,
        "violations": violations,
        "image_widths_cm": [round(shape.width.cm, 2) for shape in doc.inline_shapes],
        "image_paragraphs": image_paragraphs,
        "caption_paragraphs": caption_paragraphs,
        "table_reports": table_reports,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--manifest", required=True)
    args = parser.parse_args()

    manifest_path = Path(args.manifest)
    config = json.loads(manifest_path.read_text(encoding="utf-8"))
    master = Path(config["master"])
    output = Path(config["output"])
    chapter_dir = Path(config.get("chapter_dir", str(master.parent)))
    cache_root = Path(config.get("cache_root", str(output.parent / ".binder-cache")))
    top_headings = config.get("top_headings")
    chapters = config.get("chapters")
    outline = None
    if not top_headings or not chapters:
        outline = renumber_outline(infer_outline_from_master(master))
    if not top_headings:
        top_headings = outline_top_headings(outline or [])
    if not top_headings:
        raise ValueError("No top_headings were provided and no Heading 1 entries could be inferred from the master catalogue.")
    backup = config.get("backup", True)
    backup_strategy = config.get("backup_strategy", "once")
    renumber = config.get("renumber_headings", False)
    image_width_cm = float(config.get("image_width_cm", 16.13))
    table_profiles = config.get("table_profiles")

    output.parent.mkdir(parents=True, exist_ok=True)
    backup_path = None
    if backup and backup_strategy != "none" and output.exists():
        if backup_strategy == "timestamp":
            from datetime import datetime
            backup_path = output.with_name(f"{output.stem}_backup_{datetime.now().strftime('%Y%m%d_%H%M%S')}{output.suffix}")
            shutil.copy2(output, backup_path)
        else:
            backup_path = output.with_name(f"{output.stem}_原始备份{output.suffix}")
            if not backup_path.exists():
                shutil.copy2(output, backup_path)

    doc = Document(str(master))
    clear_body(doc)
    chapter_report = []
    if chapters:
        for chapter in chapters:
            chapter_report.append(append_chapter(doc, chapter["title"], Path(chapter["source"])))
    else:
        excluded = {master.resolve(), output.resolve()}
        candidates = candidate_files(chapter_dir, excluded, recursive=bool(config.get("recursive", True)))
        chapter_report = append_outline_nodes(doc, outline or [], candidates, cache_root, image_width_cm)
    if renumber:
        renumber_headings(doc, top_headings)
    apply_format(doc, image_width_cm=image_width_cm, table_profiles=table_profiles)
    removed_empty = clean_empty_body_paragraphs(doc)
    doc.save(str(output))

    checked = Document(str(output))
    result = audit(checked, top_headings, table_profiles=table_profiles)
    result.update({
        "output": str(output),
        "backup": str(backup_path) if backup_path else None,
        "manifest": str(manifest_path),
        "chapters": chapter_report,
        "removed_empty_paragraphs": removed_empty,
    })
    audit_path = output.with_suffix(".audit.json")
    audit_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"output": str(output), "audit": str(audit_path), "backup": str(backup_path) if backup_path else None}, ensure_ascii=False))


if __name__ == "__main__":
    main()
