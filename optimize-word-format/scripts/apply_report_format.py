from __future__ import annotations

import argparse
import json
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


HEADING_RULES = {
    "Heading 1": ("黑体", 18.0, True, WD_ALIGN_PARAGRAPH.JUSTIFY, 2.0, 6.0, 6.0, 0.0),
    "Heading 2": ("楷体", 18.0, True, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.5, None, None, 0.0),
    "Heading 3": ("仿宋", 16.0, True, WD_ALIGN_PARAGRAPH.LEFT, 1.5, None, None, 0.0),
    "Heading 4": ("仿宋", 15.0, True, WD_ALIGN_PARAGRAPH.LEFT, 1.5, None, None, 0.0),
}
NORMAL_RULE = ("仿宋", 14.0, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.5, None, None, 0.97)
TABLE_RULE = ("仿宋", 12.0, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 1.0, None, None, 0.0)
CAPTION_RULE = ("仿宋", 14.0, True, WD_ALIGN_PARAGRAPH.CENTER, 1.5, None, None, 0.0)
CAPTION_PATTERN = re.compile(r"^[图表]\s*[0-9一二三四五六七八九十]+(?:[.．、：:\-－])?\s+.+")
CONTENT_WIDTH_CM = 14.66
TABLE_PROFILES = [
    {
        "contains": ["序号", "考核指标", "完成情况", "是否已完成"],
        "widths_cm": [0.94, 2.89, 4.17, 1.74, 2.55, 2.73],
        "border": {"style": "none", "size": 0, "color": "auto"},
    },
    {
        "contains": ["指标", "市财政资金", "自筹资金", "合计"],
        "widths_cm": [4.82, 4.06, 3.08, 3.08],
        "border": {"style": "single", "size": 4, "color": "auto"},
    },
    {
        "contains": ["序号", "财政拨款资金支出科目", "计划支出金额", "已支付资金"],
        "widths_cm": [2.35, 5.93, 3.68, 3.06],
        "border": {"style": "single", "size": 6, "color": "auto"},
    },
]


def set_run_font(run, font_name: str, size_pt: float, bold: bool | None) -> None:
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold
    run.font.italic = False
    run.font.underline = False
    run.font.color.rgb = None
    rpr = run._element.get_or_add_rPr()
    rstyle = rpr.find(qn("w:rStyle"))
    if rstyle is not None:
        rpr.remove(rstyle)
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("eastAsia", "ascii", "hAnsi", "cs"):
        rfonts.set(qn(f"w:{attr}"), font_name)


def apply_style_rule(style, rule: tuple) -> None:
    font_name, size, bold, align, line_spacing, before, after, first_indent = rule
    style.font.name = font_name
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("eastAsia", "ascii", "hAnsi", "cs"):
        rfonts.set(qn(f"w:{attr}"), font_name)
    fmt = style.paragraph_format
    fmt.alignment = align
    fmt.line_spacing = line_spacing
    fmt.first_line_indent = Cm(first_indent)
    if before is not None:
        fmt.space_before = Pt(before)
    if after is not None:
        fmt.space_after = Pt(after)


def apply_paragraph_rule(paragraph, rule: tuple) -> None:
    font_name, size, bold, align, line_spacing, before, after, first_indent = rule
    fmt = paragraph.paragraph_format
    fmt.alignment = align
    fmt.line_spacing = line_spacing
    fmt.first_line_indent = Cm(first_indent)
    if before is not None:
        fmt.space_before = Pt(before)
    if after is not None:
        fmt.space_after = Pt(after)
    for run in paragraph.runs:
        if run._element.xpath(".//w:drawing"):
            continue
        set_run_font(run, font_name, size, bold)


def is_caption(text: str) -> bool:
    return bool(CAPTION_PATTERN.match((text or "").strip()))


def normalize_caption_prefix(paragraph) -> bool:
    text = paragraph.text
    stripped = text.lstrip(" \t\u3000")
    if text == stripped or not is_caption(stripped):
        return False
    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = stripped
    else:
        paragraph.add_run(stripped)
    return True


def rule_for_paragraph(paragraph) -> tuple:
    name = paragraph.style.name if paragraph.style is not None else ""
    if name in HEADING_RULES:
        return HEADING_RULES[name]
    if is_caption(paragraph.text):
        return CAPTION_RULE
    return NORMAL_RULE


def cm_to_twips(value: float) -> int:
    return int(round(value * 567))


def twips_to_cm(value: str | None) -> float | None:
    if not value:
        return None
    try:
        return round(int(value) / 567, 2)
    except ValueError:
        return None


def set_or_append(parent, tag: str):
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    return child


def set_table_width(table, width_cm: float) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = set_or_append(tbl_pr, "w:tblW")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(cm_to_twips(width_cm)))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is not None:
        tbl_pr.remove(tbl_ind)


def set_table_grid(table, widths_cm: list[float]) -> None:
    tbl_grid = table._tbl.tblGrid
    if tbl_grid is None:
        tbl_grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, tbl_grid)
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths_cm:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(cm_to_twips(float(width))))
        tbl_grid.append(col)


def set_cell_width(cell, width_cm: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = set_or_append(tc_pr, "w:tcW")
    tc_w.set(qn("w:type"), "dxa")
    tc_w.set(qn("w:w"), str(cm_to_twips(width_cm)))
    cell.width = Cm(width_cm)


def set_cell_margins(cell, top: int = 80, bottom: int = 80, left: int = 120, right: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = set_or_append(tc_pr, "w:tcMar")
    for side, value in {"top": top, "bottom": bottom, "left": left, "right": right}.items():
        node = set_or_append(tc_mar, f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, style: str = "single", size: int = 4, color: str = "auto") -> None:
    tbl_pr = table._tbl.tblPr
    borders = set_or_append(tbl_pr, "w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = set_or_append(borders, f"w:{edge}")
        node.set(qn("w:val"), style)
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def clear_table_shading(table) -> None:
    for row in table.rows:
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            shd = tc_pr.find(qn("w:shd"))
            if shd is not None:
                tc_pr.remove(shd)


def table_text(table) -> str:
    return " / ".join(cell.text.strip() for cell in table.rows[0].cells) if table.rows else ""


def table_profile(table, extra_profiles: list[dict] | None = None) -> dict:
    header = table_text(table)
    profiles = extra_profiles or []
    for profile in profiles + TABLE_PROFILES:
        if all(token in header for token in profile.get("contains", [])):
            return profile
    cols = len(table.columns)
    width = CONTENT_WIDTH_CM / cols if cols else CONTENT_WIDTH_CM
    return {"widths_cm": [width] * cols, "border": {"style": "single", "size": 4, "color": "auto"}}


def apply_table_geometry(table, extra_profiles: list[dict] | None = None) -> None:
    profile = table_profile(table, extra_profiles)
    widths = [float(w) for w in profile.get("widths_cm", [])]
    border = profile.get("border", {"style": "single", "size": 4, "color": "auto"})
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if widths:
        set_table_width(table, sum(widths))
        set_table_grid(table, widths)
    set_table_borders(table, style=border.get("style", "single"), size=int(border.get("size", 4)), color=border.get("color", "auto"))
    clear_table_shading(table)
    for row_idx, row in enumerate(table.rows):
        if row_idx == 0:
            row._tr.get_or_add_trPr()
            if row._tr.trPr.find(qn("w:tblHeader")) is None:
                row._tr.trPr.append(OxmlElement("w:tblHeader"))
        for idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if idx < len(widths):
                set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                apply_paragraph_rule(paragraph, TABLE_RULE)


def clean_empty_body_paragraphs(doc: Document) -> int:
    removed = 0
    body = doc._body._element
    for element in list(body):
        if element.tag == qn("w:sectPr"):
            continue
        if element.tag == qn("w:p") and not "".join(t.text or "" for t in element.iter(qn("w:t"))).strip() and not element.xpath(".//w:drawing"):
            body.remove(element)
            removed += 1
    return removed


def apply_report_format(doc: Document, image_width_cm: float = 16.13, table_profiles: list[dict] | None = None) -> dict:
    caption_prefix_cleaned = 0
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.17)
        section.right_margin = Cm(3.17)
        section.header_distance = Cm(1.5)
        section.footer_distance = Cm(1.75)

    for style_name, rule in HEADING_RULES.items():
        if style_name in doc.styles:
            apply_style_rule(doc.styles[style_name], rule)
    apply_style_rule(doc.styles["Normal"], NORMAL_RULE)

    for paragraph in doc.paragraphs:
        if is_caption(paragraph.text) and normalize_caption_prefix(paragraph):
            caption_prefix_cleaned += 1
        apply_paragraph_rule(paragraph, rule_for_paragraph(paragraph))
    for table in doc.tables:
        apply_table_geometry(table, table_profiles)
    for paragraph in doc.paragraphs:
        if paragraph._p.xpath(".//w:drawing"):
            paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.line_spacing = 1.5
    for shape in doc.inline_shapes:
        if shape.width:
            ratio = Cm(image_width_cm) / shape.width
            shape.width = Cm(image_width_cm)
            shape.height = int(shape.height * ratio)
    removed_empty = clean_empty_body_paragraphs(doc)
    return {"caption_prefix_cleaned": caption_prefix_cleaned, "removed_empty_paragraphs": removed_empty}


def font_name_from_run(run) -> str | None:
    rpr = run._element.rPr
    if rpr is None or rpr.rFonts is None:
        return None
    return rpr.rFonts.get(qn("w:eastAsia")) or rpr.rFonts.get(qn("w:ascii")) or rpr.rFonts.get(qn("w:hAnsi"))


def pt_value(length) -> float | None:
    return round(length.pt, 2) if length is not None else None


def table_widths(table) -> list[float | None]:
    if not table.rows:
        return []
    values = []
    for cell in table.rows[0].cells:
        tc_w = cell._tc.tcPr.find(qn("w:tcW")) if cell._tc.tcPr is not None else None
        values.append(twips_to_cm(tc_w.get(qn("w:w")) if tc_w is not None else None))
    return values


def audit(doc: Document, table_profiles: list[dict] | None = None) -> dict:
    violations = []
    captions = []
    images = []
    tables = []
    for idx, paragraph in enumerate(doc.paragraphs, 1):
        raw_text = paragraph.text
        text = raw_text.strip()
        if is_caption(raw_text):
            has_leading = raw_text != raw_text.lstrip(" \t\u3000")
            captions.append({"paragraph": idx, "text": text, "alignment": str(paragraph.paragraph_format.alignment), "has_leading_space": has_leading})
            if has_leading:
                violations.append({"issue": "caption_leading_space", "paragraph": idx, "text": raw_text[:80]})
            if paragraph.paragraph_format.alignment != WD_ALIGN_PARAGRAPH.CENTER:
                violations.append({"issue": "caption_alignment", "paragraph": idx, "actual": str(paragraph.paragraph_format.alignment)})
        if paragraph._p.xpath(".//w:drawing"):
            images.append({"paragraph": idx, "alignment": str(paragraph.paragraph_format.alignment)})
            if paragraph.paragraph_format.alignment != WD_ALIGN_PARAGRAPH.CENTER:
                violations.append({"issue": "image_paragraph_alignment", "paragraph": idx, "actual": str(paragraph.paragraph_format.alignment)})
        if not text:
            continue
        rule = rule_for_paragraph(paragraph)
        for run in paragraph.runs:
            if not run.text.strip() or run._element.xpath(".//w:drawing"):
                continue
            if font_name_from_run(run) != rule[0] or pt_value(run.font.size) != rule[1] or (rule[2] is not None and bool(run.bold) != rule[2]):
                violations.append({"issue": "run_format", "paragraph": idx, "text": text[:80]})
                break
    for idx, table in enumerate(doc.tables, 1):
        profile = table_profile(table, table_profiles)
        report = {"table": idx, "header": table_text(table), "alignment": str(table.alignment), "first_row_widths_cm": table_widths(table), "expected_widths_cm": profile.get("widths_cm", [])}
        tables.append(report)
        if table.alignment != WD_TABLE_ALIGNMENT.CENTER:
            violations.append({"issue": "table_alignment", "table": idx, "actual": str(table.alignment)})
    return {
        "paragraphs": len(doc.paragraphs),
        "tables": tables,
        "inline_shapes": len(doc.inline_shapes),
        "image_widths_cm": [round(shape.width.cm, 2) for shape in doc.inline_shapes],
        "image_paragraphs": images,
        "caption_paragraphs": captions,
        "violations": violations,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output")
    parser.add_argument("--in-place", action="store_true")
    parser.add_argument("--image-width-cm", type=float, default=16.13)
    parser.add_argument("--table-profiles-json")
    args = parser.parse_args()

    input_path = Path(args.input)
    if args.in_place:
        output_path = input_path
    elif args.output:
        output_path = Path(args.output)
    else:
        output_path = input_path.with_name(f"{input_path.stem}_格式优化{input_path.suffix}")

    table_profiles = None
    if args.table_profiles_json:
        table_profiles = json.loads(Path(args.table_profiles_json).read_text(encoding="utf-8"))

    if args.in_place and input_path.exists():
        backup = input_path.with_name(f"{input_path.stem}_格式优化前备份{input_path.suffix}")
        if not backup.exists():
            shutil.copy2(input_path, backup)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(str(input_path))
    changes = apply_report_format(doc, image_width_cm=args.image_width_cm, table_profiles=table_profiles)
    doc.save(str(output_path))

    checked = Document(str(output_path))
    result = audit(checked, table_profiles=table_profiles)
    result.update({"input": str(input_path), "output": str(output_path), **changes})
    audit_path = output_path.with_suffix(".format-audit.json")
    audit_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"output": str(output_path), "audit": str(audit_path), "violations": len(result["violations"])}, ensure_ascii=False))


if __name__ == "__main__":
    main()
